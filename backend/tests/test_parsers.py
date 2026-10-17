"""Parsers and cleaning. Fixtures are generated here: PDFs with reportlab, Word files with
python-docx, Markdown and plain text inline."""

from __future__ import annotations

import codecs
import io
import logging
import zipfile

import docx
import pytest
from docx.enum.style import WD_STYLE_TYPE
from pypdf import PageObject, PdfReader, PdfWriter
from reportlab.lib import pdfencrypt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.core.config import get_settings
from app.core.filetypes import FileType
from app.ingestion.cleaning import (
    clean_blocks,
    clean_prose,
    clean_verbatim,
    normalize_newlines,
    remove_repeated_edge_lines,
    strip_control,
    typical_line_width,
    unwrap_lines,
)
from app.ingestion.parsers import detect_file_type, parse
from app.ingestion.parsers.docx import parse_docx
from app.ingestion.parsers.markdown import parse_markdown
from app.ingestion.parsers.pdf import parse_pdf
from app.ingestion.parsers.text import decode_text, parse_text
from app.ingestion.types import Block, BlockKind, IngestionError

H, P, C, T, L = (
    BlockKind.HEADING,
    BlockKind.PARAGRAPH,
    BlockKind.CODE,
    BlockKind.TABLE,
    BlockKind.LIST,
)


def summary(blocks: list[Block]) -> list[tuple[BlockKind, str, tuple[str, ...], int | None]]:
    return [(b.kind, b.text, b.heading_path, b.page_number) for b in blocks]


def error_code(call, *args) -> str:
    with pytest.raises(IngestionError) as info:
        call(*args)
    return info.value.code


# --- PDF ----------------------------------------------------------------------------------


def make_pdf(
    pages: list[list[str]],
    *,
    header: str | None = None,
    footer: str | None = None,
    outline: list[tuple[str, int, int]] | tuple[()] = (),
    encrypt: pdfencrypt.StandardEncryption | None = None,
) -> bytes:
    """One canvas text line per string. `outline` holds (title, 0-based page, level)
    bookmarks in reading order; `footer` may use {page} and {pages}."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, encrypt=encrypt)
    for number, lines in enumerate(pages):
        pdf.setFont("Helvetica", 10)
        if header:
            pdf.drawString(72, 800, header)
        for i, (title, page, level) in enumerate(outline):
            if page == number:
                pdf.bookmarkPage(f"b{i}")
                pdf.addOutlineEntry(title, f"b{i}", level=level)
        for i, line in enumerate(lines):
            pdf.drawString(72, 760 - 14 * i, line)
        if footer:
            pdf.drawString(72, 40, footer.format(page=number + 1, pages=len(pages)))
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


REPORT = [
    [
        "Introduction",
        "The quarterly report covers revenue, costs and the outlook for the inter-",
        "national markets in which the company operates.",
        "Growth was strong in every region.",
    ],
    [
        "Methods",
        "We surveyed customers in every region and compared the answers with the",
        "results of the previous year to measure the change in their satisfaction.",
        "Sampling",
        "Respondents were chosen at random from the customer database of the firm",
        "with equal weight for each region.",
    ],
    [
        "Respondents in the final week answered by phone because the web form was",
        "unavailable.",
    ],
]
OUTLINE = [("Introduction", 0, 0), ("Methods", 1, 0), ("Sampling", 1, 1)]


def report_pdf(**kwargs) -> bytes:
    kwargs.setdefault("outline", OUTLINE)
    return make_pdf(REPORT, header="ACME Confidential", footer="Page {page} of {pages}", **kwargs)


def test_pdf_pages_sections_headers_and_wrapped_lines():
    parsed = parse_pdf(report_pdf())
    assert parsed.page_count == 3
    assert parsed.stats == {"pages_failed": 0, "pages_without_text": 0}
    methods, sampling = ("Methods",), ("Methods", "Sampling")
    assert summary(parsed.blocks) == [
        (H, "Introduction", ("Introduction",), 1),
        (
            P,
            "The quarterly report covers revenue, costs and the outlook for the international "
            "markets in which the company operates.\nGrowth was strong in every region.",
            ("Introduction",),
            1,
        ),
        (H, "Methods", methods, 2),
        (
            P,
            "We surveyed customers in every region and compared the answers with the results "
            "of the previous year to measure the change in their satisfaction.",
            methods,
            2,
        ),
        # Found mid-page: the heading line starts the subsection and is not glued to the
        # paragraph above it.
        (H, "Sampling", sampling, 2),
        (
            P,
            "Respondents were chosen at random from the customer database of the firm with "
            "equal weight for each region.",
            sampling,
            2,
        ),
        # No bookmark on page 3: the section continues.
        (
            P,
            "Respondents in the final week answered by phone because the web form was unavailable.",
            sampling,
            3,
        ),
    ]
    text = "\n".join(block.text for block in parsed.blocks)
    assert "ACME" not in text and "Page" not in text  # running header and footer


def test_pdf_without_outline_has_paragraphs_only():
    parsed = parse_pdf(report_pdf(outline=()))
    # Extraction yields no blank lines here, so each page is one paragraph block.
    assert [(b.kind, b.heading_path, b.page_number) for b in parsed.blocks] == [
        (P, (), 1),
        (P, (), 2),
        (P, (), 3),
    ]
    assert parsed.blocks[1].text.startswith("Methods\nWe surveyed customers")


def test_pdf_bookmark_whose_title_is_not_on_the_page_starts_at_the_top():
    parsed = parse_pdf(report_pdf(outline=[*OUTLINE, ("Phone interviews", 2, 1)]))
    last = parsed.blocks[-1]
    assert (last.kind, last.heading_path, last.page_number) == (
        P,
        ("Methods", "Phone interviews"),
        3,
    )
    assert "Phone interviews" not in [block.text for block in parsed.blocks]


def test_pdf_short_documents_keep_repeated_lines():
    # Two pages are too few to tell a running header from content.
    parsed = parse_pdf(make_pdf(REPORT[:2], header="ACME Confidential"))
    assert parsed.blocks[0].text.startswith("ACME Confidential\n")


def test_pdf_with_user_password_is_rejected():
    encryption = pdfencrypt.StandardEncryption("secret", ownerPassword="owner")
    data = make_pdf([["Launch plans for the new product."]], encrypt=encryption)
    assert error_code(parse_pdf, data) == "encrypted_pdf"


def test_pdf_with_owner_password_only_is_read():
    encryption = pdfencrypt.StandardEncryption("", ownerPassword="owner", canPrint=0)
    parsed = parse_pdf(make_pdf([["Launch plans for the new product."]], encrypt=encryption))
    assert [block.text for block in parsed.blocks] == ["Launch plans for the new product."]


def test_aes_pdf_with_owner_password_only_is_read():
    # AES needs the `cryptography` package (pypdf[crypto]); RC4 above does not.
    writer = PdfWriter(clone_from=io.BytesIO(make_pdf([["Launch plans for the new product."]])))
    writer.encrypt(user_password="", owner_password="owner", algorithm="AES-256")
    buffer = io.BytesIO()
    writer.write(buffer)
    parsed = parse_pdf(buffer.getvalue())
    assert [block.text for block in parsed.blocks] == ["Launch plans for the new product."]


def test_pdf_that_cannot_be_decrypted_counts_as_encrypted(monkeypatch):
    def unsupported(self, password):
        raise NotImplementedError("AES requires the cryptography package")

    monkeypatch.setattr(PdfReader, "decrypt", unsupported)
    encryption = pdfencrypt.StandardEncryption("", ownerPassword="owner")
    data = make_pdf([["Launch plans."]], encrypt=encryption)
    assert error_code(parse_pdf, data) == "encrypted_pdf"


def test_pdf_without_text_layer_is_rejected():
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.rect(100, 100, 300, 300, fill=1)  # a "scanned" page: graphics only
    pdf.showPage()
    pdf.save()
    with pytest.raises(IngestionError) as info:
        parse_pdf(buffer.getvalue())
    assert info.value.code == "no_extractable_text"
    assert "OCR" in info.value.message
    # Page numbers alone are not text either.
    assert error_code(parse_pdf, make_pdf([["1"], ["2"], ["3"]])) == "no_extractable_text"


def test_pdf_with_too_many_pages_is_rejected(monkeypatch):
    monkeypatch.setattr(get_settings(), "max_pdf_pages", 2)
    with pytest.raises(IngestionError) as info:
        parse_pdf(report_pdf())
    assert info.value.code == "too_many_pages"
    assert "3 pages" in info.value.message


@pytest.mark.parametrize(
    "data",
    [b"", b"definitely not a pdf", b"%PDF-1.7\n%%EOF", report_pdf()[:1500]],
    ids=["empty", "garbage", "header-only", "truncated"],
)
def test_corrupt_pdf_is_rejected(data):
    assert error_code(parse_pdf, data) == "corrupt_file"


def test_pdf_page_that_fails_to_extract_is_skipped(monkeypatch):
    original = PageObject.extract_text

    def flaky(self, *args, **kwargs):
        text = original(self, *args, **kwargs)
        if "Methods" in text:
            raise KeyError("/Font")
        return text

    monkeypatch.setattr(PageObject, "extract_text", flaky)
    parsed = parse_pdf(report_pdf())
    assert parsed.stats == {"pages_failed": 1, "pages_without_text": 1}
    assert "surveyed" not in "".join(block.text for block in parsed.blocks)
    assert parsed.blocks[-1].page_number == 3

    monkeypatch.setattr(PageObject, "extract_text", lambda self, *a, **k: 1 / 0)
    assert error_code(parse_pdf, report_pdf()) == "corrupt_file"


def test_pdf_unreadable_outline_costs_only_the_section_titles(monkeypatch, caplog):
    def broken(self):
        raise ValueError("bad outline")

    monkeypatch.setattr(PdfReader, "outline", property(broken))
    with caplog.at_level(logging.WARNING, logger="app.ingestion.parsers.pdf"):
        parsed = parse_pdf(report_pdf())
    assert {block.heading_path for block in parsed.blocks} == {()}
    assert "ingestion.pdf_outline_unreadable" in caplog.messages


# --- DOCX ---------------------------------------------------------------------------------


def make_docx() -> bytes:
    document = docx.Document()
    document.add_paragraph("Employee Handbook", style="Title")
    document.add_paragraph("Welcome to the company.")
    document.add_heading("Benefits", level=1)
    document.add_paragraph("All employees receive benefits.")
    document.add_paragraph("Health insurance", style="List Bullet")
    document.add_paragraph("Dental plan included", style="List Bullet 2")
    document.add_paragraph("Pension", style="List Number")
    numbered = document.add_paragraph("Numbered child")  # Word numbering, nesting level 2
    num_pr = numbered._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 2
    num_pr.get_or_add_numId().val = 1
    document.add_heading("Costs", level=2)
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Cost | monthly"
    table.cell(1, 0).text = "Basic"
    table.cell(1, 1).text = "$10\nper month"
    nested = table.cell(1, 0).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "nested a"
    nested.cell(0, 1).text = "nested | b"  # row 3 stays empty and is dropped
    document.add_heading("Tools", level=1)
    document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("def greet(name):", style="Code")
    document.add_paragraph("    return f'hi {name}'", style="Code")
    document.add_paragraph("Run it like this:")
    document.add_paragraph().add_run("greet('ada')").font.name = "Courier New"
    document.add_paragraph("", style="Code")
    document.add_paragraph("The end.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_headings_lists_tables_and_code():
    parsed = parse_docx(make_docx())
    assert parsed.page_count is None
    root = ("Employee Handbook",)
    benefits, costs, tools = (*root, "Benefits"), (*root, "Benefits", "Costs"), (*root, "Tools")
    assert summary(parsed.blocks) == [
        (H, "Employee Handbook", root, None),
        (P, "Welcome to the company.", root, None),
        (H, "Benefits", benefits, None),
        (P, "All employees receive benefits.", benefits, None),
        (
            L,
            "- Health insurance\n  - Dental plan included\n- Pension\n    - Numbered child",
            benefits,
            None,
        ),
        (H, "Costs", costs, None),
        (
            T,
            "| Plan | Cost \\| monthly |\n| --- | --- |\n"
            "| Basic nested a nested \\| b | $10 per month |",
            costs,
            None,
        ),
        (H, "Tools", tools, None),
        (C, "```\ndef greet(name):\n    return f'hi {name}'\n```", tools, None),
        (P, "Run it like this:", tools, None),
        (C, "```\ngreet('ada')\n```", tools, None),
        (P, "The end.", tools, None),
    ]
    # Cleaning keeps list nesting and code indentation.
    cleaned = clean_blocks(parsed.blocks)
    assert cleaned[4].text == parsed.blocks[4].text
    assert cleaned[8].text == parsed.blocks[8].text


def test_corrupt_docx_is_rejected():
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zipped:
        zipped.writestr("hello.txt", "not a word document")
    assert error_code(parse_docx, b"PK\x03\x04 broken") == "corrupt_file"
    assert error_code(parse_docx, archive.getvalue()) == "corrupt_file"
    assert error_code(parse_docx, make_docx()[:2000]) == "corrupt_file"


def test_docx_that_unpacks_too_large_is_rejected(monkeypatch):
    monkeypatch.setattr(get_settings(), "max_docx_uncompressed_mb", 0)
    assert error_code(parse_docx, make_docx()) == "file_too_large"


# --- Markdown -----------------------------------------------------------------------------

MARKDOWN = """---
title: Engineering Guide
owner: platform
---
# Engineering Guide

Intro paragraph that wraps
onto a second line.  \nHard break kept.\\
Also kept.

Setext Section
--------------

Run the server:

    make run
    make test

## Build *and* [Deploy](https://example.com/deploy) ##

```python
# not a heading
def main():
    print("|not|a|table|")
```

~~~
# also not a heading
~~~

| Command | Purpose |
|---------|:-------:|
| `make test` | run tests |

- first item
  continued line
- second item
  - nested item
1. numbered

***

Other Top
=========

Text after a thematic break.

```bash
echo "unclosed fence"
"""


def test_markdown_structure():
    guide = ("Engineering Guide",)
    setext, build = (*guide, "Setext Section"), (*guide, "Build and Deploy")
    assert summary(parse_markdown(MARKDOWN).blocks) == [
        (H, "Engineering Guide", guide, None),
        (
            P,
            "Intro paragraph that wraps onto a second line.\nHard break kept.\nAlso kept.",
            guide,
            None,
        ),
        (H, "Setext Section", setext, None),
        (P, "Run the server:", setext, None),
        (C, "    make run\n    make test", setext, None),
        (H, "Build and Deploy", build, None),
        (
            C,
            '```python\n# not a heading\ndef main():\n    print("|not|a|table|")\n```',
            build,
            None,
        ),
        (C, "~~~\n# also not a heading\n~~~", build, None),
        (
            T,
            "| Command | Purpose |\n|---------|:-------:|\n| `make test` | run tests |",
            build,
            None,
        ),
        (
            L,
            "- first item\n  continued line\n- second item\n  - nested item\n1. numbered",
            build,
            None,
        ),
        (H, "Other Top", ("Other Top",), None),
        (P, "Text after a thematic break.", ("Other Top",), None),
        (C, '```bash\necho "unclosed fence"\n```', ("Other Top",), None),
    ]


def test_markdown_heading_levels_and_edge_cases():
    text = "# A\n\n### C\n\nunder c\n\n## B\n\nunder b\n\n#\n\n#hashtag is text\n"
    blocks = [b for b in parse_markdown(text).blocks if b.kind is not H]
    assert [(b.text, b.heading_path) for b in blocks] == [
        ("under c", ("A", "C")),
        ("under b", ("A", "B")),
        ("#hashtag is text", ("A", "B")),
    ]


def test_markdown_without_closing_front_matter_keeps_the_text():
    blocks = parse_markdown("---\nnot front matter\n").blocks
    assert summary(blocks) == [(P, "not front matter", (), None)]


# --- Plain text ---------------------------------------------------------------------------


def test_text_paragraphs_and_preformatted_blocks():
    text = "First paragraph\nstill first.\n\n   \n    indented\n    block\n\nLast one."
    assert summary(parse_text(text).blocks) == [
        (P, "First paragraph\nstill first.", (), None),
        (C, "    indented\n    block", (), None),
        (P, "Last one.", (), None),
    ]


def test_decode_text_encodings():
    cafe = "caf" + chr(0xE9)
    assert decode_text(codecs.BOM_UTF8 + f"{cafe}\r\nline".encode()) == f"{cafe}\nline"
    assert decode_text(f"{cafe}\n".encode("utf-16")) == f"{cafe}\n"  # BOM + native order
    assert decode_text(codecs.BOM_UTF16_BE + cafe.encode("utf-16-be")) == cafe
    assert decode_text(cafe.encode("cp1252")) == cafe  # legacy Windows text
    assert decode_text(b"plain\rold mac") == "plain\nold mac"
    # 0x81 and 0x8D are neither UTF-8 nor Windows-1252.
    assert error_code(decode_text, b"\x81\x8d") == "unreadable_text"


def test_parse_dispatches_on_file_type():
    assert parse(b"# Title\n\nBody", FileType.MD).blocks[0].kind is H
    assert parse(b"# Title\n\nBody", FileType.TXT).blocks[0].text == "# Title"
    assert parse(make_docx(), FileType.DOCX).blocks[0].text == "Employee Handbook"
    assert parse(report_pdf(), FileType.PDF).page_count == 3


def test_detect_file_type():
    assert detect_file_type("text/markdown", "notes.txt") is FileType.MD  # MIME wins
    assert detect_file_type("application/octet-stream", "Guide.MD") is FileType.MD
    assert detect_file_type("", "notes.markdown") is FileType.MD
    assert detect_file_type("application/pdf", "x") is FileType.PDF
    assert error_code(detect_file_type, "application/zip", "archive.zip") == (
        "unsupported_file_type"
    )


# --- Cleaning -----------------------------------------------------------------------------


def test_control_characters_and_newlines():
    zero_width, bom, line_separator = chr(0x200B), chr(0xFEFF), chr(0x2028)
    assert strip_control(f"a\x00b{zero_width}c{bom}\td\x07") == "abc\td"
    assert normalize_newlines(f"a\r\nb\rc{line_separator}d\fe") == "a\nb\nc\nd\n\ne"


def test_clean_prose_normalizes_unicode_and_whitespace():
    ligature, wide_a, nbsp = chr(0xFB01), chr(0xFF21), chr(0xA0)
    text = f"  {ligature}nal   {wide_a}{nbsp}team \n\n\n\nnext\x00 line  "
    assert clean_prose(text) == "final A team\n\nnext line"
    assert clean_prose("- a\n  - b   c", keep_indent=True) == "- a\n  - b c"
    assert clean_prose("- a\n  - b") == "- a\n- b"


def test_clean_verbatim_keeps_layout():
    ligature = chr(0xFB01)
    code = f"\n\ndef f():  \n    return '{ligature}'\x00\n\n\n    # end\n\n"
    assert clean_verbatim(code) == f"def f():\n    return '{ligature}'\n\n\n    # end"


def test_clean_blocks():
    blocks = [
        Block(P, f"Hello{chr(0xA0)}  world", ("  Guide ", ""), 1),
        Block(C, "  x = 1\n\n  y = 2  ", ("Guide",)),
        Block(T, "| a |  b |\n| --- | --- |"),
        Block(P, f" \x00 {chr(0x200B)} "),  # nothing left: dropped
    ]
    assert summary(clean_blocks(blocks)) == [
        (P, "Hello world", ("Guide",), 1),
        (C, "  x = 1\n\n  y = 2", ("Guide",), None),
        (T, "| a |  b |\n| --- | --- |", (), None),
    ]


def test_remove_repeated_edge_lines():
    words = ["alpha", "beta", "gamma", "delta"]
    pages = [
        ["Annual Report 2025", f"Body text about {word}.", "Summary", f"Page {n}"]
        for n, word in enumerate(words, start=1)
    ]
    cleaned = remove_repeated_edge_lines(pages)
    # "Summary" is among the last three lines and repeats, so it goes too; body survives.
    assert cleaned == [[f"Body text about {word}."] for word in words]
    pages[1] = ["Different first page", *pages[1][1:]]
    assert remove_repeated_edge_lines(pages)[1][0] == "Different first page"
    assert remove_repeated_edge_lines(pages[:2]) == pages[:2]  # too few pages to tell


def test_remove_repeated_edge_lines_keeps_repeats_in_the_middle():
    middle = ["one", "two", "three", "Repeated mid-page line", "four", "five", "six"]
    pages = [[f"{line} {n}" if line != middle[3] else line for line in middle] for n in range(3)]
    assert all(middle[3] in page for page in remove_repeated_edge_lines(pages))


def test_typical_line_width_is_the_90th_percentile():
    assert typical_line_width([["x" * n for n in range(1, 11)], [""]]) == 9
    assert typical_line_width([]) == 0


def test_unwrap_lines():
    width = 40
    long_a = "This sentence is long enough to wrap at"  # 39 characters
    lines = [
        long_a,
        "the column edge.",
        "Short line.",
        "",
        "A hyphenated word ends this line: inter-",
        "national, and Hewlett-Packard is a name-",
        "Company. Eventually a line breaks after",
        "- a list item starts a new line",
    ]
    assert unwrap_lines(lines, width) == [
        f"{long_a} the column edge.",
        "Short line.",
        "",
        "A hyphenated word ends this line: international, and Hewlett-Packard is a "
        "name-Company. Eventually a line breaks after",
        "- a list item starts a new line",
    ]
