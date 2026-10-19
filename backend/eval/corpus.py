"""The synthetic corpus: a fictional company, its people, workspaces and 18 documents.

Sources live in eval/sources/. Markdown and text documents are uploaded as written; PDF
and DOCX documents are generated here from a Markdown source (`*.pdf.md`, `*.docx.md`)
with reportlab and python-docx, so page numbers, PDF outlines, Word heading styles and
Word tables go through the real parsers. Builds are deterministic: the same source always
gives the same bytes (reportlab's invariant mode, fixed DOCX timestamps).

The access rule in `can_read` restates docs/ARCHITECTURE.md independently of
app/services/access.py, so the evaluation can check the application against it.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from app.ingestion.parsers.markdown import parse_markdown
from app.ingestion.types import BlockKind
from app.models import OrgRole, WorkspaceRole

SOURCES = Path(__file__).parent / "sources"
ORG_NAME = "Northwind Aerial Robotics (fictional)"
EMAIL_DOMAIN = "demo.example.com"
SYNTHETIC_MARKER = "SYNTHETIC DOCUMENT"
UPLOADER = "manager"  # every document is uploaded by the knowledge manager
# Chunking the evaluation pins (the product defaults), so needles map to chunks the same
# way in tests/test_eval_dataset.py and in the evaluated index.
CHUNKING = {"target_tokens": 350, "max_tokens": 500, "overlap_tokens": 50}

_FIXED_TIME = datetime(2026, 1, 1)
_ZIP_TIME = (2026, 1, 1, 0, 0, 0)
_LIST_MARKER = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+")
_TABLE_DELIMITER = re.compile(r"^\|?(?:\s*:?-+:?\s*\|)+\s*:?-*:?\s*$")


@dataclass(frozen=True)
class Person:
    key: str
    full_name: str
    org_role: OrgRole
    workspaces: dict[str, WorkspaceRole] = field(default_factory=dict)

    @property
    def email(self) -> str:
        return f"{self.key}@{EMAIL_DOMAIN}"


@dataclass(frozen=True)
class WorkspaceSpec:
    key: str
    name: str
    description: str


@dataclass(frozen=True)
class Doc:
    key: str
    source: str  # file name in eval/sources/
    title: str
    workspace: str
    restricted: bool = False
    grants: tuple[str, ...] = ()  # people with an explicit read grant

    @property
    def format(self) -> str:
        """pdf, docx, md or txt: `x.pdf.md` is the Markdown source of a PDF."""
        parts = self.source.split(".")
        return parts[-2] if len(parts) == 3 else parts[-1]

    @property
    def filename(self) -> str:
        return f"{self.source.split('.')[0]}.{self.format}"

    def source_text(self) -> str:
        return (SOURCES / self.source).read_text(encoding="utf-8")


WORKSPACES = (
    WorkspaceSpec("people", "People & Policies", "Company policies, benefits and finance."),
    WorkspaceSpec("engineering", "Engineering", "API reference, runbooks, postmortems, firmware."),
    WorkspaceSpec("customer", "Customer Success", "Product FAQ, pricing and SLA, maintenance."),
)

_ALL = {ws.key: WorkspaceRole.VIEWER for ws in WORKSPACES}
PEOPLE = (
    Person("admin", "Demo Admin", OrgRole.ORG_ADMIN),
    Person(
        "manager",
        "Demo Knowledge Manager",
        OrgRole.KNOWLEDGE_MANAGER,
        {ws.key: WorkspaceRole.OWNER for ws in WORKSPACES},
    ),
    Person("employee", "Demo Employee", OrgRole.EMPLOYEE, dict(_ALL)),
    Person("finance", "Demo Finance Lead", OrgRole.EMPLOYEE, {"people": WorkspaceRole.VIEWER}),
    Person(
        "contractor", "Demo Contractor", OrgRole.EMPLOYEE, {"engineering": WorkspaceRole.VIEWER}
    ),
)

DOCUMENTS = (
    Doc("travel-2025", "travel-policy-2025.pdf.md", "Travel Policy 2025", "people"),
    Doc("travel-2026", "travel-policy-2026.pdf.md", "Travel Policy 2026", "people"),
    Doc("expenses", "expense-policy.docx.md", "Expense Reimbursement Policy", "people"),
    Doc("benefits-uk", "hr-benefits-uk.md", "UK Employee Benefits Guide", "people"),
    Doc("benefits-us", "hr-benefits-us.md", "US Employee Benefits Guide", "people"),
    Doc("infosec", "information-security-policy.pdf.md", "Information Security Policy", "people"),
    Doc("retention", "data-retention-schedule.docx.md", "Data Retention Schedule", "people"),
    Doc("remote-work", "remote-work-policy.txt", "Remote Work Policy", "people"),
    Doc("conduct", "code-of-conduct.txt", "Code of Business Conduct", "people"),
    Doc(
        "suppliers", "supplier-onboarding-checklist.txt", "Supplier Onboarding Checklist", "people"
    ),
    Doc(
        "exec-comp",
        "executive-compensation-2026.docx.md",
        "Executive Compensation and Project Albatross 2026",
        "people",
        restricted=True,
        grants=("finance",),
    ),
    Doc("api-reference", "skyfleet-api-reference.md", "SkyFleet API Reference v3", "engineering"),
    Doc(
        "runbook", "telemetry-ingest-runbook.md", "Runbook: Telemetry Ingest Service", "engineering"
    ),
    Doc(
        "postmortem",
        "incident-postmortem-inc-2291.pdf.md",
        "Postmortem INC-2291: Telemetry Upload Outage",
        "engineering",
    ),
    Doc("firmware", "firmware-release-notes.md", "Kestrel Firmware Release Notes", "engineering"),
    Doc("product-faq", "product-faq.md", "Kestrel X2 Product FAQ", "customer"),
    Doc(
        "pricing-sla",
        "pricing-and-sla.docx.md",
        "SkyFleet Pricing and Service Level Agreement",
        "customer",
    ),
    Doc(
        "maintenance",
        "drone-maintenance-guide.md",
        "Kestrel X2 Field Maintenance Guide",
        "customer",
    ),
)

PEOPLE_BY_KEY = {person.key: person for person in PEOPLE}
DOCUMENTS_BY_KEY = {doc.key: doc for doc in DOCUMENTS}
WORKSPACES_BY_KEY = {ws.key: ws for ws in WORKSPACES}


def can_read(person: Person, doc: Doc) -> bool:
    """The documented read rule, the same for org admins: a workspace-visible document is
    readable by its workspace's members, and a restricted one only by people with a user
    grant, its uploader (while a member) and the workspace owners."""
    if person.key in doc.grants:
        return True
    role = person.workspaces.get(doc.workspace)
    if role is None:
        return False
    return not doc.restricted or role is WorkspaceRole.OWNER or person.key == UPLOADER


def build(doc: Doc) -> bytes:
    """The file as it is uploaded."""
    text = doc.source_text()
    if doc.format == "pdf":
        return build_pdf(text, doc.title)
    if doc.format == "docx":
        return build_docx(text, doc.title)
    return text.encode("utf-8")


def build_pdf(text: str, title: str) -> bytes:
    """A4 pages: every level-2 section starts a new page, every heading is a bookmark
    (outline entry) and every page has a running footer the cleaner should remove."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas

    page_width, page_height = A4
    margin = 72.0
    text_width = page_width - 2 * margin
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, invariant=1)
    pdf.setTitle(title)
    pdf.setAuthor(ORG_NAME)
    state = {"y": page_height - margin, "page": 1}

    def finish_page() -> None:
        pdf.setFont("Helvetica", 8)
        pdf.drawString(margin, 40, f"{ORG_NAME} | {title} | Page {state['page']}")
        pdf.showPage()
        state["page"] += 1
        state["y"] = page_height - margin

    def draw(line: str, font: str, size: float, indent: float = 0.0) -> None:
        if state["y"] < margin + size:
            finish_page()
        pdf.setFont(font, size)
        pdf.drawString(margin + indent, state["y"], line)
        state["y"] -= size * 1.45

    sections = 0
    for number, block in enumerate(parse_markdown(text).blocks):
        if block.kind is BlockKind.HEADING:
            level = len(block.heading_path)
            if level == 2:
                sections += 1
                if sections > 1:
                    finish_page()
            state["y"] -= 6
            pdf.bookmarkPage(f"h{number}")
            pdf.addOutlineEntry(block.text, f"h{number}", level=level - 1)
            draw(block.text, "Helvetica-Bold", 16 if level == 1 else 12)
            state["y"] -= 4
        elif block.kind is BlockKind.PARAGRAPH:
            for line in simpleSplit(block.text, "Helvetica", 10, text_width):
                draw(line, "Helvetica", 10)
            state["y"] -= 6
        elif block.kind is BlockKind.LIST:
            for item in block.text.split("\n"):
                lines = simpleSplit(item.strip(), "Helvetica", 10, text_width - 12)
                for index, line in enumerate(lines):
                    draw(line, "Helvetica", 10, indent=12 if index else 0)
            state["y"] -= 6
        else:
            raise ValueError(f"{block.kind.value} blocks are not supported in PDF sources")
    finish_page()
    pdf.save()
    return buffer.getvalue()


def build_docx(text: str, title: str) -> bytes:
    """Title and Heading N styles, List Bullet paragraphs and real Word tables."""
    import docx

    document = docx.Document()
    properties = document.core_properties
    properties.title = title
    properties.author = properties.last_modified_by = ORG_NAME
    properties.created = properties.modified = _FIXED_TIME
    properties.revision = 1
    for block in parse_markdown(text).blocks:
        if block.kind is BlockKind.HEADING:
            document.add_heading(block.text, level=len(block.heading_path) - 1)
        elif block.kind is BlockKind.PARAGRAPH:
            document.add_paragraph(block.text)
        elif block.kind is BlockKind.LIST:
            for item in block.text.split("\n"):
                document.add_paragraph(_LIST_MARKER.sub("", item), style="List Bullet")
        elif block.kind is BlockKind.TABLE:
            rows = [cells for line in block.text.split("\n") if (cells := _cells(line)) is not None]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            raise ValueError(f"{block.kind.value} blocks are not supported in DOCX sources")
    buffer = io.BytesIO()
    document.save(buffer)
    return _fixed_zip_times(buffer.getvalue())


def _cells(line: str) -> list[str] | None:
    """Cells of a Markdown table row; None for the delimiter row."""
    if _TABLE_DELIMITER.match(line.strip()):
        return None
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _fixed_zip_times(data: bytes) -> bytes:
    """python-docx stamps zip entries with the current time; fix them for stable bytes."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=_ZIP_TIME)
            target.writestr(fixed, source.read(info), compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()
